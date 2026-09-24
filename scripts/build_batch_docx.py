# -*- coding: utf-8 -*-
"""
build_batch_docx.py — Phase 6.

Batch = angle x awareness level, numbered B1..B15 in a 5x3 round
(B1 = angle 1 level A, B2 = angle 1 level B, B4 = angle 2 level A). Images: B1-V1, B1-V2, B1-V3.

Naming carries the funnel stage: {AUTHOR} {BRAND-SKU} {TEST}-{FUNNEL}-B{n}-V{k}
  e.g. AA BRAND-SKU T101-TF-B1-V1   (TF = top of funnel, FF = bottom of funnel). Always hyphens.

(1) Builds two .docx per batch, in the house document standard (Arial 12pt, 1.5 line
    spacing, black), both in the batch folder next to the 3 images:
      AA BRAND-SKU T101-TF-B1/AA BRAND-SKU T101-TF-B1.docx
          main document, each field with the label in bold and the content on the line
          below: ANGLE (name + one-line summary), AWARENESS LEVEL (label + one-line
          summary), PROFILE, HEADLINE, DESCRIPTION, CTA, COPY
      AA BRAND-SKU T101-TF-B1/AA BRAND-SKU T101-TF-B1 - PROMPTS.docx
          the real prompt of each image, read from image-brief.md
(2) Builds the round upload sheet (upload.csv), one row per ad, with the profile
    column and the classification columns that creative-intel reads afterwards.

Frontmatter fields read from the draft:
    angle, angle_num, level               (as before)
    angle_summary                         the angle in one sentence
    level_summary                         optional; defaults to the level's standard line
    profile                               the Meta page that runs the batch
CTA by funnel stage: TF = Learn more, FF = Shop now (--funnel decides, not the draft).
Awareness levels (default): A Problem aware, B Solution aware, C Hidden cause.

Usage:
    python build_batch_docx.py --draft drafts/B1.md   (writes to AA BRAND-SKU T101-TF-B1/)
    python build_batch_docx.py --folder drafts/ --test T101 --product BRAND-SKU --funnel TF
    python build_batch_docx.py --sheet --folder drafts/ --test T101 --product BRAND-SKU --funnel TF \
           --destination "https://..." --out upload.csv
"""
import argparse, csv, re, sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLACK = RGBColor(0, 0, 0)
# fallback only: the label of each variation comes from the image brief, because the
# image variable (ethnicity, age band, composition) changes from round to round
ETHNICITY = {"1": "variation 1", "2": "variation 2", "3": "variation 3"}
DEFAULT_AUTHOR = "AA"

# ------------------------------------------------------------------ parsing
def parse_front(txt):
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", txt, re.S)
    if not m:
        return {}, txt
    fm, body = {}, txt[m.end():]
    for line in m.group(1).split("\n"):
        if ":" not in line:
            continue
        k, v = line.split(":", 1)
        v = v.strip()
        if v.startswith("[") and v.endswith("]"):
            v = [x.strip() for x in v[1:-1].split(",") if x.strip()]
        fm[k.strip()] = v
    return fm, body

def section(body, *titles):
    for t in titles:
        m = re.search(r"^##\s*%s\s*$" % re.escape(t), body, re.M | re.I)
        if m:
            rest = body[m.end():]
            n = re.search(r"^##\s+", rest, re.M)
            return (rest[:n.start()] if n else rest).strip()
    return ""

# ------------------------------------------------------------------ image brief
LEVELS = {"A": "Problem aware", "B": "Solution aware", "C": "Hidden cause"}
LEVEL_SUMMARY = {
    "A": "Opens on the scene, the mechanism arrives late.",
    "B": "Opens on what he already tried, the mechanism explains why it failed.",
    "C": "Opens on the mechanism reveal, the pain comes after.",
}
FUNNELS = {"TF": "Top of funnel", "FF": "Bottom of funnel"}
CTA_BY_FUNNEL = {"TF": "Learn more", "FF": "Shop now"}

def read_brief(cell, path="image-brief.md"):
    """Locked scene of the cell + the label of each variation (from its #### header)."""
    p = Path(path)
    if not p.exists():
        return None
    txt = p.read_text(encoding="utf-8")
    # the lookahead must exclude '####', otherwise the block closes at the first variation
    pat_cell = r"^###[^#\n][^\n]*\b" + re.escape(cell) + r"\b(.*?)(?=^###[^#]|\Z)"
    m = re.search(pat_cell, txt, re.M | re.S)
    if not m:
        return None
    block = m.group(0)
    scene = ""
    ms = re.search(r"\*\*Locked scene:?\*\*\s*([^\n]+)", block)
    if ms:
        scene = ms.group(1).strip()
    labels = {}
    # '.' does not match newlines without DOTALL, so this only takes the header line
    pat_lab = r"^####.*?\b" + re.escape(cell) + r"-[vV](\d)\b(.*)$"
    for ml in re.finditer(pat_lab, block, re.M):
        labels[ml.group(1)] = ml.group(2).strip().lstrip("-" + chr(8212) + chr(8211) + " ").strip()
    return {"scene": scene, "labels": labels}

def read_prompts(cell, path="image-brief.md"):
    """Full prompt of each variation of the cell, keyed by '1', '2', '3'."""
    p = Path(path)
    if not p.exists():
        return {}
    txt = p.read_text(encoding="utf-8")
    out = {}
    pat = r"^####[^\n]*\b" + re.escape(cell) + r"-[vV](\d)\b[^\n]*\n(.*?)(?=^#{3,4}\s|\Z)"
    for m in re.finditer(pat, txt, re.M | re.S):
        pm = re.search(r"\*\*Prompt:?\*\*\s*\n(.*)", m.group(2), re.S)
        out[m.group(1)] = (pm.group(1) if pm else m.group(2)).strip()
    return out

def base_name(author, product, test, funnel, cell):
    """AA BRAND-SKU T101-TF-B1: author, product, test, funnel stage and batch, always hyphens."""
    return "%s %s %s-%s-%s" % (author, product, test, funnel, cell)

def cell_path(author, product, test, cell, funnel="TF"):
    """Each batch has its own folder holding the two docx and the 3 images:
    AA BRAND-SKU T101-TF-B1/AA BRAND-SKU T101-TF-B1.docx (batch = angle x level, B1..B15)"""
    name = base_name(author, product, test, funnel, cell)
    Path(name).mkdir(exist_ok=True)
    return str(Path(name) / (name + ".docx"))

def batch_order(p):
    """B2 before B10: sort by batch number, not alphabetically."""
    m = re.match(r"B(\d+)$", Path(p).stem)
    return (0, int(m.group(1))) if m else (1, Path(p).stem)

# ------------------------------------------------------------------ docx
def _run(run, size=12, bold=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(attr), "Arial")

def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.top_margin, s.left_margin = Cm(3), Cm(3)
    s.bottom_margin, s.right_margin = Cm(2), Cm(2)
    n = doc.styles["Normal"]
    n.font.name = "Arial"; n.font.size = Pt(12); n.font.color.rgb = BLACK
    pf = n.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_after = Pt(10)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def heading(doc, text, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    _run(p.add_run(text.upper()), 12, bold=True)
    return p

def body_text(doc, text, justify=True):
    for block in [b for b in text.split("\n") if b.strip()]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(10)
        _run(p.add_run(block.strip()), 12, bold=False)

def field(doc, label, text, justify=True):
    """Label in bold on one line, content on the line below."""
    heading(doc, label)
    body_text(doc, text, justify=justify)

def hyphens(text):
    """Separators are always hyphens, never a middle dot."""
    return text.replace(" · ", " - ")

def clean(text):
    return "\n".join(l for l in text.splitlines() if l.strip() != "---").strip()

def angle_label(fm):
    n, name = fm.get("angle_num", ""), fm.get("angle", "")
    return ("Angle %s - %s" % (n, name)) if n else name

def build_docx(md_path, out_path, brief="image-brief.md", funnel="TF"):
    fm, body = parse_front(Path(md_path).read_text(encoding="utf-8"))
    primary = section(body, "PRIMARY TEXT")
    base = Path(out_path).stem
    # batch = angle x level, numbered B1..B15; the level comes from the frontmatter
    mcell = re.search(r"-(B\d+)$", base)
    cell = mcell.group(1) if mcell else ""
    level = str(fm.get("level", ""))
    cta = CTA_BY_FUNNEL[funnel]
    cta_md = section(body, "CTA")
    if cta_md and cta_md.lower() != cta.lower():
        print("  warning %s: draft says CTA '%s', funnel %s requires '%s'" % (cell, cta_md, funnel, cta))
    for k in ("profile", "angle_summary"):
        if not fm.get(k):
            print("  warning %s: frontmatter has no '%s'" % (cell, k))

    # main document: everything whoever uploads the ad needs, except the prompts
    doc = new_doc()
    heading(doc, base, center=True)
    field(doc, "Angle", "\n".join(x for x in (angle_label(fm), fm.get("angle_summary", "")) if x))
    field(doc, "Awareness level",
          "\n".join((LEVELS.get(level, level), fm.get("level_summary") or LEVEL_SUMMARY.get(level, ""))))
    field(doc, "Profile", fm.get("profile") or "PENDING")
    field(doc, "Headline", section(body, "LINK HEADLINE"))
    field(doc, "Description", section(body, "LINK DESCRIPTION"))
    field(doc, "CTA", cta)
    field(doc, "Copy", primary)
    doc.save(out_path)

    # prompts document, separate, in the same folder: {standard name} - PROMPTS.docx
    prompts = read_prompts(cell, brief) if cell else {}
    labels = (read_brief(cell, brief) or {}).get("labels", {})
    pr = new_doc()
    heading(pr, base + " - PROMPTS", center=True)
    for k in ("1", "2", "3"):
        cab = pr.add_paragraph()
        cab.paragraph_format.space_before = Pt(14)
        cab.paragraph_format.space_after = Pt(6)
        _run(cab.add_run("V%s - %s-V%s.png" % (k, base, k)), 12, bold=True)
        body_text(pr, hyphens(labels.get(k) or ETHNICITY[k]), justify=False)
        body_text(pr, clean(prompts.get(k, "")), justify=False)
    pr.save(Path(out_path).parent / (base + " - PROMPTS.docx"))
    return out_path, len(primary)

# ------------------------------------------------------------------ sheet
def build_sheet(folder, test, product, destination, campaign, out, author=DEFAULT_AUTHOR,
                funnel="TF", brief="image-brief.md"):
    rows = []
    for md in sorted(Path(folder).glob("*.md"), key=batch_order):
        fm, body = parse_front(md.read_text(encoding="utf-8"))
        m = re.match(r"B(\d+)$", md.stem)
        if not m:
            continue
        b, level = m.group(1), str(fm.get("level", ""))
        base = base_name(author, product, test, funnel, "B%s" % b)
        labels = (read_brief("B%s" % b, brief) or {}).get("labels", {})
        primary = section(body, "PRIMARY TEXT")
        head = section(body, "LINK HEADLINE")
        desc = section(body, "LINK DESCRIPTION")
        cta = CTA_BY_FUNNEL[funnel]
        for v in ("1", "2", "3"):
            rows.append({
                "campaign": campaign or "%s Long Form %s" % (test, funnel),
                "ad_set": "%s-%s-B%s" % (test, funnel, b),
                "profile": fm.get("profile", ""),
                "ad": "%s-V%s" % (base, v),
                "image": "%s-V%s.png" % (base, v),
                "primary_text": primary,
                "headline": head,
                "description": desc,
                "cta": cta,
                "destination": destination or "",
                "angle": fm.get("angle", ""),
                "angle_num": fm.get("angle_num", ""),
                "level": "%s (%s)" % (LEVELS.get(level, level), level),
                "variation": "V%s" % v,
                "cluster": labels.get(v) or ETHNICITY[v],
                "chars": len(primary),
            })
    if not rows:
        print("No valid draft found in %s" % folder); return
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()), delimiter=";")
        w.writeheader(); w.writerows(rows)
    print("upload sheet: %s  (%d ads, %d ad sets)"
          % (out, len(rows), len({r['ad_set'] for r in rows})))

# ------------------------------------------------------------------ cli
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--draft"); ap.add_argument("--folder")
    ap.add_argument("--out"); ap.add_argument("--sheet", action="store_true")
    ap.add_argument("--test", default="T101"); ap.add_argument("--product", default="BRAND-SKU")
    ap.add_argument("--author", default=DEFAULT_AUTHOR)
    ap.add_argument("--funnel", default="TF", choices=sorted(FUNNELS),
                    help="TF top of funnel (CTA Learn more), FF bottom of funnel (CTA Shop now)")
    ap.add_argument("--brief", default="image-brief.md")
    ap.add_argument("--destination", default=""); ap.add_argument("--campaign", default="")
    a = ap.parse_args()

    if a.sheet:
        build_sheet(a.folder or "drafts", a.test, a.product, a.destination,
                    a.campaign, a.out or "upload.csv", a.author, a.funnel, a.brief)
        return
    if a.draft:
        out = a.out or cell_path(a.author, a.product, a.test, Path(a.draft).stem, a.funnel)
        out, n = build_docx(a.draft, out, a.brief, a.funnel)
        print("built: %s  (%d chars of copy)" % (out, n)); return
    if a.folder:
        for md in sorted(Path(a.folder).glob("*.md"), key=batch_order):
            if not re.match(r"B\d+$", md.stem):
                continue
            out = cell_path(a.author, a.product, a.test, md.stem, a.funnel)
            _, n = build_docx(md, out, a.brief, a.funnel)
            print("built: %s  (%d chars)" % (out, n))
        return
    ap.error("pass --draft, --folder or --sheet")

if __name__ == "__main__":
    main()
